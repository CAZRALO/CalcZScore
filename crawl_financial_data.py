#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
=============================================================================
CRAWLER BÁO CÁO TÀI CHÍNH 700 CÔNG TY NIÊM YẾT SÀN HOSE & HNX
Dự án: Hệ thống phân tích Altman Z-Score & Z''-Score (CalcZScore)
Nguồn danh sách: Sheet 'DANH MỤC' trong file 'dữ_liệu_2 2.xlsx' (700 mã)
Nguồn dữ liệu BCTC: CafeF Báo cáo tài chính & SSI iBoard
Các chỉ tiêu trích xuất: TSNH, NNH, TTS, LNCPP, LNTT, CPLV, VHTT/VCSH, TNPT, DTT
Giai đoạn: 2021 - 2025
=============================================================================
"""

import os
import sys
import re
import json
import time
import glob
import random
import argparse
import datetime
from concurrent.futures import ThreadPoolExecutor, as_completed

import requests
from bs4 import BeautifulSoup
import openpyxl

# Thiết lập UTF-8 cho console Windows
if sys.platform == 'win32':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
        sys.stderr.reconfigure(encoding='utf-8')
    except Exception:
        pass

# =============================================================================
# CẤU HÌNH & HẰNG SỐ
# =============================================================================
DEFAULT_WORKERS = 6
DEFAULT_YEARS = [2021, 2022, 2023, 2024, 2025]
CHECKPOINT_FILE = "crawl_checkpoint.json"
USER_AGENTS = [
    'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/123.0.0.0 Safari/537.36',
    'Mozilla/5.0 (Windows NT 10.0; Win64; x64; rv:124.0) Gecko/20100101 Firefox/124.0',
    'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/123.0.0.0 Safari/537.36',
    'Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/123.0.0.0 Safari/537.36'
]

BANK_TICKERS = {
    'VCB', 'BID', 'CTG', 'TCB', 'MBB', 'ACB', 'VPB', 'STB', 'HDB', 'SHB', 'TPB',
    'LPB', 'MSB', 'OCB', 'VIB', 'SSB', 'EIB', 'BAB', 'NAB', 'BVB', 'KLB', 'PGB',
    'SGB', 'ABB', 'VAB', 'SSI', 'VND', 'HCM', 'VCI', 'SHS', 'MBS', 'CTS', 'BSI',
    'FTS', 'AGR', 'VIX', 'ORS', 'TVS', 'APG', 'WSS', 'IVS', 'BVH', 'BMI', 'PVI',
    'MIG', 'BIC', 'PTI', 'PRE', 'VNR'
}

# =============================================================================
# HÀM BỔ TRỢ XỬ LÝ SỐ LIỆU & BCTC
# =============================================================================
def get_headers():
    return {
        'User-Agent': random.choice(USER_AGENTS),
        'Accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,image/webp,*/*;q=0.8',
        'Accept-Language': 'vi-VN,vi;q=0.9,en-US;q=0.8,en;q=0.7'
    }

def clean_val(s):
    """Làm sạch chuỗi số từ bảng CafeF và quy đổi ra đơn vị TỶ ĐỒNG (làm tròn)."""
    if not s:
        return 0
    s = s.strip()
    is_neg = s.startswith('-') or (s.startswith('(') and s.endswith(')'))
    s = re.sub(r'[^\d.,]', '', s)
    if not s:
        return 0
    try:
        val_str = s.replace('.', '').replace(',', '.')
        num = float(val_str)
        val_in_billions = round(num / 1e9)
        return -abs(val_in_billions) if is_neg else val_in_billions
    except Exception:
        return 0

def parse_years(soup):
    """Phát hiện các năm hiển thị trên dòng tiêu đề bảng BCTC CafeF."""
    detected = []
    for tr in soup.find_all('tr'):
        cells = [td.get_text(strip=True) for td in tr.find_all(['th', 'td'])]
        cands = [c for c in cells if re.match(r'^\d{4}$', c) and 2000 <= int(c) <= 2030]
        if len(cands) >= 2:
            return cands
    return detected

def extract_metric(soup, row_years, match_fn):
    """Trích xuất giá trị của một chỉ tiêu theo hàm nhận diện tiêu đề."""
    res = {}
    for tr in soup.find_all('tr'):
        cells = [td.get_text(strip=True) for td in tr.find_all(['th', 'td'])]
        if len(cells) > len(row_years):
            row_title = cells[0].upper()
            if match_fn(row_title):
                for idx, yr in enumerate(row_years):
                    res[yr] = clean_val(cells[idx + 1])
                return res
    return res

# =============================================================================
# HÀM CRAWL TỪNG DOANH NGHIỆP
# =============================================================================
def crawl_single_stock(item, target_years=DEFAULT_YEARS, max_retries=2):
    """
    Crawl toàn bộ BCTC giai đoạn target_years và giá thị trường cho một doanh nghiệp.
    item: dict chứa {'stt': int, 'symbol': str, 'name': str, 'exchange': str}
    """
    ticker = item['symbol'].strip().upper()
    is_financial = ticker in BANK_TICKERS

    result = {
        'stt': item['stt'],
        'symbol': ticker,
        'name': item.get('name', ticker),
        'exchange': item.get('exchange', ''),
        'price': None,
        'is_financial': is_financial,
        'status': 'pending',
        'years_found': 0,
        'data': {
            'tsnh': {}, 'nnh': {}, 'tts': {}, 'lncpp': {},
            'lntt': {}, 'cplv': {}, 'vhtt': {}, 'tnpt': {}, 'dtt': {}
        }
    }

    # 1. Trích xuất Giá thị trường và Tên tiếng Việt từ SSI iBoard
    for attempt in range(max_retries):
        try:
            ssi_url = f"https://iboard-query.ssi.com.vn/stock/{ticker}"
            r = requests.get(ssi_url, headers=get_headers(), timeout=7)
            if r.status_code == 200:
                ssi_data = r.json().get('data', {})
                if ssi_data:
                    raw_price = ssi_data.get('refPrice') or ssi_data.get('priorClosePrice') or 0
                    if raw_price > 0:
                        result['price'] = round(raw_price / 1000.0, 2)
                    if ssi_data.get('companyNameVi'):
                        result['name'] = ssi_data['companyNameVi']
                    if ssi_data.get('exchange'):
                        result['exchange'] = ssi_data['exchange'].upper()
                    elif ssi_data.get('market') == 'UPX':
                        result['exchange'] = 'UPCOM'
                break
        except Exception:
            time.sleep(0.5)

    # Nếu là Ngân hàng/Tổ chức tài chính, mô hình Altman không áp dụng
    if is_financial:
        result['status'] = 'financial_institution_skipped'
        return result

    # 2. Trích xuất BCTC từ CafeF
    # Target endpoint years: [2025, 2024, 2021] đủ phủ các năm 2018 - 2025
    query_years = [2025, 2024, 2021]
    
    for yr in query_years:
        for attempt in range(max_retries):
            try:
                bs_url = f"https://s.cafef.vn/bao-cao-tai-chinh/{ticker}/bsheet/{yr}/0/0/0/bao-cao.chn"
                inc_url = f"https://s.cafef.vn/bao-cao-tai-chinh/{ticker}/incsta/{yr}/0/0/0/bao-cao.chn"
                
                headers = get_headers()
                r_bs = requests.get(bs_url, headers=headers, timeout=8)
                r_inc = requests.get(inc_url, headers=headers, timeout=8)

                if r_bs.status_code != 200 and r_inc.status_code != 200:
                    continue

                soup_bs = BeautifulSoup(r_bs.text, 'html.parser') if r_bs.status_code == 200 else None
                soup_inc = BeautifulSoup(r_inc.text, 'html.parser') if r_inc.status_code == 200 else None

                row_years = []
                if soup_bs:
                    row_years = parse_years(soup_bs)
                if not row_years and soup_inc:
                    row_years = parse_years(soup_inc)

                if not row_years:
                    break

                # Trích xuất các chỉ tiêu Bảng Cân Đối Kế Toán
                if soup_bs:
                    tsnh = extract_metric(soup_bs, row_years, lambda n: 'TÀI SẢN NGẮN HẠN' in n and any(x in n for x in ['A-', 'A.', 'A -']))
                    tts  = extract_metric(soup_bs, row_years, lambda n: 'TỔNG CỘNG TÀI SẢN' in n or ('TỔNG TÀI SẢN' in n and 'DÀI HẠN' not in n))
                    nnh  = extract_metric(soup_bs, row_years, lambda n: 'NỢ NGẮN HẠN' in n and any(x in n for x in ['I.', 'I -', 'I-']))
                    tnpt = extract_metric(soup_bs, row_years, lambda n: 'NỢ PHẢI TRẢ' in n and any(x in n for x in ['C.', 'C -', 'C-']))
                    vcsh = extract_metric(soup_bs, row_years, lambda n: 'VỐN CHỦ SỞ HỮU' in n and any(x in n for x in ['D.', 'D -', 'D-']))
                    lncpp= extract_metric(soup_bs, row_years, lambda n: 'CHƯA PHÂN PHỐI' in n)

                    for y_str in row_years:
                        y = int(y_str)
                        if y in target_years:
                            if y_str in tsnh and y_str not in result['data']['tsnh']: result['data']['tsnh'][y_str] = tsnh[y_str]
                            if y_str in nnh and y_str not in result['data']['nnh']: result['data']['nnh'][y_str] = nnh[y_str]
                            if y_str in tts and y_str not in result['data']['tts']: result['data']['tts'][y_str] = tts[y_str]
                            if y_str in lncpp and y_str not in result['data']['lncpp']: result['data']['lncpp'][y_str] = lncpp[y_str]
                            if y_str in vcsh and y_str not in result['data']['vhtt']: result['data']['vhtt'][y_str] = vcsh[y_str]
                            if y_str in tnpt and y_str not in result['data']['tnpt']: result['data']['tnpt'][y_str] = tnpt[y_str]

                # Trích xuất các chỉ tiêu Báo Cáo Kết Quả Kinh Doanh
                if soup_inc:
                    dtt  = extract_metric(soup_inc, row_years, lambda n: 'DOANH THU THUẦN' in n)
                    lntt = extract_metric(soup_inc, row_years, lambda n: 'LỢI NHUẬN KẾ TOÁN TRƯỚC THUẾ' in n or 'LỢI NHUẬN TRƯỚC THUẾ' in n)
                    cplv = extract_metric(soup_inc, row_years, lambda n: 'CHI PHÍ LÃI VAY' in n)

                    for y_str in row_years:
                        y = int(y_str)
                        if y in target_years:
                            if y_str in dtt and y_str not in result['data']['dtt']: result['data']['dtt'][y_str] = dtt[y_str]
                            if y_str in lntt and y_str not in result['data']['lntt']: result['data']['lntt'][y_str] = lntt[y_str]
                            if y_str in cplv and y_str not in result['data']['cplv']: result['data']['cplv'][y_str] = cplv[y_str]

                break
            except Exception:
                time.sleep(0.5)

    # 3. Tính toán bổ sung tính nhất quán Bảng Cân Đối (TTS = VCSH + TNPT)
    years_with_data = 0
    for y in target_years:
        ys = str(y)
        # Bổ sung tính nhất quán
        if not result['data']['tts'].get(ys) and result['data']['vhtt'].get(ys) and result['data']['tnpt'].get(ys):
            result['data']['tts'][ys] = result['data']['vhtt'][ys] + result['data']['tnpt'][ys]
        if not result['data']['tnpt'].get(ys) and result['data']['tts'].get(ys) and result['data']['vhtt'].get(ys):
            result['data']['tnpt'][ys] = max(0, result['data']['tts'][ys] - result['data']['vhtt'][ys])
        if not result['data']['vhtt'].get(ys) and result['data']['tts'].get(ys) and result['data']['tnpt'].get(ys):
            result['data']['vhtt'][ys] = max(0, result['data']['tts'][ys] - result['data']['tnpt'][ys])

        # Kiểm tra năm có dữ liệu thực tế
        has_tts = result['data']['tts'].get(ys, 0) > 0
        has_dtt = result['data']['dtt'].get(ys, 0) > 0 or result['data']['lntt'].get(ys, 0) != 0
        if has_tts or has_dtt:
            years_with_data += 1

    result['years_found'] = years_with_data
    result['status'] = 'success' if years_with_data > 0 else 'no_data'
    return result

# =============================================================================
# HÀM ĐỌC DANH SÁCH TỪ FILE EXCEL
# =============================================================================
def find_excel_file(custom_path=None):
    """Tìm file excel chứa danh sách doanh nghiệp."""
    if custom_path and os.path.exists(custom_path):
        return custom_path
    
    # Tìm kiếm các file xlsx trong thư mục hiện tại
    candidates = glob.glob("*.xlsx")
    for f in candidates:
        if "dữ_liệu" in f.lower() or "du_lieu" in f.lower() or "dữ_liệu" in f.lower():
            return f
    if candidates:
        return candidates[0]
    return None

def load_stock_list_from_excel(excel_path):
    """Đọc sheet 'DANH MỤC' để lấy 700 doanh nghiệp."""
    wb = openpyxl.load_workbook(excel_path, read_only=True)
    if 'DANH MỤC' not in wb.sheetnames:
        raise ValueError(f"Không tìm thấy sheet 'DANH MỤC' trong file: {excel_path}")
    
    ws = wb['DANH MỤC']
    stocks = []
    
    for row in ws.iter_rows(min_row=2, max_col=6, values_only=True):
        stt, ticker, name, exchange, price, num_years = row[:6]
        if stt is not None and ticker:
            stocks.append({
                'stt': int(stt),
                'symbol': str(ticker).strip().upper(),
                'name': str(name).strip() if name else str(ticker).strip().upper(),
                'exchange': str(exchange).strip() if exchange else '',
                'price': float(price) if price is not None else None,
                'existing_years': int(num_years) if num_years is not None else 0
            })
    
    wb.close()
    return stocks

# =============================================================================
# HÀM CẬP NHẬT GHI VÀO FILE EXCEL
# =============================================================================
def update_excel_with_crawled_data(excel_path, output_path, crawl_results, target_years=DEFAULT_YEARS):
    """
    Cập nhật dữ liệu đã crawl vào sheet 'DỮ LIỆU' và 'DANH MỤC' của workbook.
    Giữ nguyên toàn bộ cấu trúc định dạng ban đầu của file Excel.
    """
    print(f"\n📝 Đang nạp workbook: {excel_path}...")
    wb = openpyxl.load_workbook(excel_path)
    
    ws_dl = wb['DỮ LIỆU']
    ws_dm = wb['DANH MỤC']
    
    # 1. Xác định vị trí dòng tiêu đề của từng năm trong sheet DỮ LIỆU
    year_row_headers = {}
    for r in range(1, 3600):
        val = ws_dl.cell(r, 1).value
        try:
            if int(val) in target_years:
                year_row_headers[int(val)] = r
        except (ValueError, TypeError):
            pass

    print(f"📌 Đã xác định các khối năm trong sheet 'DỮ LIỆU': {year_row_headers}")

    # Bản đồ thứ tự cột trong DỮ LIỆU:
    # 1: STT, 2: MÃ DN, 3: TSNH, 4: NNH, 5: TTS, 6: LNCPP, 7: LNTT, 8: CPLV, 9: VHTT/VCSH, 10: TNPT, 11: DTT
    metric_cols = {
        'tsnh': 3,
        'nnh': 4,
        'tts': 5,
        'lncpp': 6,
        'lntt': 7,
        'cplv': 8,
        'vhtt': 9,
        'tnpt': 10,
        'dtt': 11
    }

    updated_data_cells = 0
    updated_catalog_stocks = 0

    # Lập chỉ mục kết quả crawl theo mã chứng khoán
    results_by_ticker = {r['symbol']: r for r in crawl_results}

    # 2. Cập nhật sheet DỮ LIỆU
    for yr in target_years:
        header_row = year_row_headers.get(yr)
        if not header_row:
            continue
        
        ys = str(yr)
        for stt in range(1, 701):
            row_idx = header_row + stt
            cell_ticker = ws_dl.cell(row_idx, 2).value
            if not cell_ticker:
                continue
            
            ticker = str(cell_ticker).strip().upper()
            crawled = results_by_ticker.get(ticker)
            if not crawled or crawled.get('status') != 'success':
                continue

            cdata = crawled.get('data', {})
            # Kiểm tra xem năm này có dữ liệu không
            tts_val = cdata.get('tts', {}).get(ys)
            dtt_val = cdata.get('dtt', {}).get(ys)
            if tts_val or dtt_val:
                for metric, col_idx in metric_cols.items():
                    val = cdata.get(metric, {}).get(ys)
                    if val is not None:
                        ws_dl.cell(row_idx, col_idx).value = val
                        updated_data_cells += 1

    # 3. Cập nhật sheet DANH MỤC (GIÁ và SỐ NĂM ĐÃ CÓ SỐ LIỆU)
    for r in range(2, ws_dm.max_row + 1):
        cell_ticker = ws_dm.cell(r, 2).value
        if not cell_ticker:
            continue
        
        ticker = str(cell_ticker).strip().upper()
        crawled = results_by_ticker.get(ticker)
        if not crawled:
            continue

        # Cập nhật giá thị trường nếu lấy được từ sàn
        if crawled.get('price') is not None and crawled.get('price') > 0:
            ws_dm.cell(r, 5).value = crawled['price']

        # Cập nhật số năm có số liệu
        years_found = crawled.get('years_found', 0)
        if years_found > 0:
            existing = ws_dm.cell(r, 6).value or 0
            ws_dm.cell(r, 6).value = max(int(existing), years_found)
            updated_catalog_stocks += 1

    # 4. Ghi thêm nhật ký vào sheet GHI CHÚ
    if 'GHI CHÚ' in wb.sheetnames:
        ws_gc = wb['GHI CHÚ']
        now_str = datetime.datetime.now().strftime("%d/%m/%Y %H:%M:%S")
        log_msg = (
            f"[{now_str}] Cập nhật tự động bởi Auto-Crawler 700 DN (HOSE & HNX): "
            f"Đã xử lý {len(crawl_results)} mã, điền {updated_data_cells} ô số liệu BCTC (2021-2025) "
            f"và cập nhật trạng thái {updated_catalog_stocks} mã trong DANH MỤC."
        )
        ws_gc.append([log_msg])

    print(f"💾 Đang lưu file Excel kết quả ra: {output_path}...")
    wb.save(output_path)
    wb.close()
    print(f"✅ Đã ghi thành công {updated_data_cells} ô số liệu vào {output_path}!")

# =============================================================================
# HÀM ĐỒNG BỘ VỚI ỨNG DỤNG WEB (STOCK_DIRECTORY.JS)
# =============================================================================
def sync_with_web_app(crawl_results, all_stocks, stock_directory_path="stock_directory.js"):
    """
    Đồng bộ toàn bộ danh bạ 700 doanh nghiệp vào file stock_directory.js và xuất file
    cơ sở dữ liệu crawled_700_database.json cho Web App.
    """
    # 1. Lưu file JSON cơ sở dữ liệu đầy đủ cho web nạp tức thì
    json_path = "crawled_700_database.json"
    cleaned_dict = {}
    for r in crawl_results:
        cleaned_dict[r['symbol']] = {
            'symbol': r['symbol'],
            'companyName': r['name'],
            'exchange': r['exchange'],
            'price': r['price'],
            'years_found': r['years_found'],
            'is_financial': r['is_financial'],
            'data': r['data']
        }

    with open(json_path, 'w', encoding='utf-8') as f:
        json.dump(cleaned_dict, f, ensure_ascii=False, indent=2)
    print(f"✅ Đã xuất cơ sở dữ liệu JSON cho Web App: {json_path}")

    # 2. Cập nhật danh bạ STOCK_DIRECTORY trong stock_directory.js nếu tồn tại
    if not os.path.exists(stock_directory_path):
        return

    try:
        with open(stock_directory_path, 'r', encoding='utf-8') as f:
            js_content = f.read()

        # Nhận diện khối khai báo const STOCK_DIRECTORY = [...]
        pattern = r'const STOCK_DIRECTORY = \[(.*?)\];\s*// Helper'
        m = re.search(pattern, js_content, re.DOTALL)
        if not m:
            pattern = r'const STOCK_DIRECTORY = \[(.*?)\];'
            m = re.search(pattern, js_content, re.DOTALL)

        if m:
            existing_dict = {}
            for line in m.group(1).split('\n'):
                line_clean = line.strip()
                if line_clean.startswith('{') and (line_clean.endswith('},') or line_clean.endswith('}')):
                    s_m = re.search(r"s:\s*'([^']+)'", line_clean)
                    n_m = re.search(r"n:\s*'([^']+)'", line_clean)
                    e_m = re.search(r"e:\s*'([^']+)'", line_clean)
                    ind_m = re.search(r"ind:\s*'([^']+)'", line_clean)
                    if s_m:
                        sym = s_m.group(1).upper()
                        existing_dict[sym] = {
                            's': sym,
                            'n': n_m.group(1) if n_m else '',
                            'e': e_m.group(1) if e_m else '',
                            'ind': ind_m.group(1) if ind_m else ''
                        }

            # Tạo danh sách hợp nhất 700 mã
            merged_list = []
            for es in all_stocks:
                sym = es['symbol'].strip().upper()
                exch = 'HOSE' if es.get('exchange') == 'HSX' else (es.get('exchange') or 'HOSE')
                name = es.get('name') or sym
                
                if sym in existing_dict:
                    ex = existing_dict[sym]
                    merged_list.append({
                        's': sym,
                        'n': ex['n'] or name,
                        'e': ex['e'] or exch,
                        'ind': ex['ind'] or ''
                    })
                else:
                    merged_list.append({
                        's': sym,
                        'n': name,
                        'e': exch,
                        'ind': ''
                    })

            # Tạo lại khối JavaScript
            lines = ["const STOCK_DIRECTORY = ["]
            for item in merged_list:
                escaped_n = item['n'].replace("'", "\\'")
                escaped_ind = item['ind'].replace("'", "\\'")
                lines.append(f"  {{ s: '{item['s']}', n: '{escaped_n}', e: '{item['e']}', ind: '{escaped_ind}' }},")
            lines.append("];\n\n// Helper")

            new_js = js_content[:m.start()] + "\n".join(lines) + js_content[m.end():]
            with open(stock_directory_path, 'w', encoding='utf-8') as f:
                f.write(new_js)

            print(f"✅ Đã đồng bộ {len(merged_list)} mã vào danh bạ tìm kiếm STOCK_DIRECTORY ({stock_directory_path})!")
    except Exception as e:
        print(f"⚠️ Không thể cập nhật stock_directory.js: {e}")

# =============================================================================
# CHƯƠNG TRÌNH CHÍNH (CLI RUNNER)
# =============================================================================
def main():
    parser = argparse.ArgumentParser(description="Tool crawl BCTC 700 công ty HOSE & HNX từ file dữ_liệu_2 2.xlsx")
    parser.add_argument('--input', '-i', type=str, default=None, help="Đường dẫn file Excel gốc")
    parser.add_argument('--output', '-o', type=str, default=None, help="Đường dẫn file Excel xuất kết quả")
    parser.add_argument('--in-place', action='store_true', help="Ghi đè trực tiếp lên file Excel gốc")
    parser.add_argument('--workers', '-w', type=int, default=DEFAULT_WORKERS, help="Số luồng tải song song (mặc định: 6)")
    parser.add_argument('--limit', '-l', type=int, default=None, help="Giới hạn số lượng mã cần crawl (để test nhanh)")
    parser.add_argument('--all', '-a', action='store_true', help="Crawl toàn bộ danh sách 700 doanh nghiệp")
    parser.add_argument('--missing', action='store_true', help="Chỉ crawl các mã hiện chưa có đủ số liệu trong file gốc")
    parser.add_argument('--symbols', '-s', type=str, default=None, help="Crawl các mã cụ thể (cách nhau bởi dấu phẩy, ví dụ: AAA,HPG,FPT)")
    parser.add_argument('--resume', action='store_true', help="Tiếp tục từ checkpoint đã lưu trước đó")
    parser.add_argument('--no-sync-web', action='store_true', help="Không cập nhật file database cho Web App")
    
    args = parser.parse_args()

    excel_file = find_excel_file(args.input)
    if not excel_file:
        print("❌ Lỗi: Không tìm thấy file Excel dữ liệu trong thư mục!")
        sys.exit(1)

    output_file = excel_file if args.in_place else (args.output or "dữ_liệu_700_crawled.xlsx")

    print("=" * 70)
    print("🚀 HỆ THỐNG CRAWLER DỮ LIỆU TÀI CHÍNH 700 CÔNG TY (HOSE & HNX)")
    print(f"📂 File nguồn danh mục : {excel_file}")
    print(f"📂 File kết quả xuất ra: {output_file}")
    print(f"⚡ Số luồng xử lý     : {args.workers}")
    print("=" * 70)

    # Đọc danh sách doanh nghiệp từ Excel
    all_stocks = load_stock_list_from_excel(excel_file)
    print(f"📋 Tổng số mã tìm thấy trong sheet 'DANH MỤC': {len(all_stocks)}")

    # Lọc danh sách theo tùy chọn CLI
    target_stocks = all_stocks
    if args.symbols:
        requested = [s.strip().upper() for s in args.symbols.split(',') if s.strip()]
        target_stocks = [s for s in target_stocks if s['symbol'] in requested]
        print(f"🎯 Lọc theo mã yêu cầu: {len(target_stocks)} mã ({', '.join(requested)})")
    elif args.missing:
        target_stocks = [s for s in target_stocks if s['existing_years'] < 5]
        print(f"🎯 Lọc các mã chưa có đủ 5 năm số liệu: {len(target_stocks)} mã")

    if args.limit and args.limit > 0:
        target_stocks = target_stocks[:args.limit]
        print(f"⏱️ Giới hạn chạy thử nghiệm: {len(target_stocks)} mã đầu tiên")

    # Xử lý checkpoint nạp lại nếu có
    completed_cache = {}
    if args.resume and os.path.exists(CHECKPOINT_FILE):
        try:
            with open(CHECKPOINT_FILE, 'r', encoding='utf-8') as f:
                completed_cache = json.load(f)
            print(f"🔄 Nạp lại checkpoint thành công: Đã có sẵn {len(completed_cache)} mã.")
        except Exception as e:
            print(f"⚠️ Không thể đọc file checkpoint: {e}")

    # Lọc các mã chưa hoàn thành
    queue = [s for s in target_stocks if s['symbol'] not in completed_cache]
    print(f"⏳ Cần thực hiện crawl: {len(queue)} mã (Đã có sẵn: {len(target_stocks) - len(queue)} mã)")

    results = list(completed_cache.values())
    total_to_crawl = len(queue)
    completed_count = 0
    start_time = time.time()

    if total_to_crawl > 0:
        print("\n🚀 Đang tiến hành crawl dữ liệu song song...")
        with ThreadPoolExecutor(max_workers=args.workers) as executor:
            future_to_stock = {executor.submit(crawl_single_stock, s): s for s in queue}
            
            for future in as_completed(future_to_stock):
                stock = future_to_stock[future]
                ticker = stock['symbol']
                try:
                    res = future.result()
                    results.append(res)
                    completed_cache[ticker] = res
                    completed_count += 1

                    # Hiển thị tiến độ
                    status_icon = "✅" if res['status'] == 'success' else ("⚠️" if res['is_financial'] else "⚪")
                    elapsed = time.time() - start_time
                    speed = completed_count / elapsed if elapsed > 0 else 0
                    eta = (total_to_crawl - completed_count) / speed if speed > 0 else 0
                    
                    print(
                        f"[{completed_count:3d}/{total_to_crawl}] {status_icon} {ticker:<5} | "
                        f"Giá: {res['price'] if res['price'] else '--':>6} | "
                        f"Số năm: {res['years_found']}/5 | "
                        f"Tốc độ: {speed:.1f} mã/s | ETA: {eta:.0f}s"
                    )

                    # Lưu checkpoint mỗi 10 mã
                    if completed_count % 10 == 0:
                        with open(CHECKPOINT_FILE, 'w', encoding='utf-8') as f:
                            json.dump(completed_cache, f, ensure_ascii=False)

                except Exception as e:
                    print(f"❌ Lỗi khi crawl {ticker}: {e}")

        # Lưu checkpoint cuối cùng
        with open(CHECKPOINT_FILE, 'w', encoding='utf-8') as f:
            json.dump(completed_cache, f, ensure_ascii=False)

    print("\n" + "=" * 70)
    print("🎉 HOÀN THÀNH QUÁ TRÌNH CRAWL DỮ LIỆU!")
    total_success = sum(1 for r in results if r.get('status') == 'success')
    total_fin = sum(1 for r in results if r.get('is_financial'))
    print(f"📊 Tổng số mã hoàn thành      : {len(results)}")
    print(f"✅ Doanh nghiệp có số liệu BCTC: {total_success}")
    print(f"⚠️ Tổ chức tài chính (bỏ qua) : {total_fin}")
    print(f"⏱️ Tổng thời gian thực thi     : {time.time() - start_time:.1f} giây")
    print("=" * 70)

    # Cập nhật kết quả vào Excel
    update_excel_with_crawled_data(excel_file, output_file, results)

    # Đồng bộ với ứng dụng web nếu không bị tắt
    if not args.no_sync_web:
        sync_with_web_app(results, all_stocks)

    print("\n✨ Hoàn tất xuất sắc! Dữ liệu đã sẵn sàng để nghiên cứu và phân tích.")

if __name__ == '__main__':
    main()
